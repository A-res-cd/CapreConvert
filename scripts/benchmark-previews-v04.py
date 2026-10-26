"""Measure reopen versus validation-pixel preview paths and native document previews."""

import argparse
import importlib.util
import json
import multiprocessing
import platform
import sys
import tempfile
import time
from contextlib import closing
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def probe(source, output, reuse, connection):
    from PIL import Image

    from localraw.memory import process_memory
    from localraw.raster import image_derivatives

    started = time.perf_counter()
    with Image.open(source) as image:
        image.load()  # output validation always required
        if reuse:
            image_derivatives(image, Path(output), {"min_free_bytes": 0})
    if not reuse:
        with Image.open(source) as image:
            image.load()
            image_derivatives(image, Path(output), {"min_free_bytes": 0})
    memory = process_memory(__import__("os").getpid())
    connection.send(
        {
            "wall_seconds": time.perf_counter() - started,
            "worker_peak_bytes": memory.peak_rss,
            "reuse_validated_pixels": reuse,
        }
    )
    connection.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=ROOT / "test-results/v04-preview-benchmark.json")
    args = parser.parse_args()
    if args.output.exists():
        parser.error("Report exists; choose a new output")
    from docx import Document
    from PIL import Image

    from localraw.config import Settings
    from localraw.jobs import Jobs
    from localraw.models import JobRequest, UploadRequest
    from localraw.previews import prepare
    from localraw.storage import Workspace, file_hash, opaque_id
    from localraw.uploads import Uploads

    spec = importlib.util.spec_from_file_location(
        "memory_benchmark", ROOT / "scripts/benchmark-memory.py"
    )
    benchmark = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(benchmark)
    spec = importlib.util.spec_from_file_location("package_smoke", ROOT / "scripts/smoke-package.py")
    smoke = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(smoke)
    report = {"environment": platform.platform(), "preview_probe": [], "documents": []}
    with tempfile.TemporaryDirectory(prefix="preview-v04-", dir=ROOT / "test-results") as temporary:
        directory = Path(temporary)
        with Image.open(ROOT / "test-results/memory-optimized/real/output.jpg") as image:
            image.thumbnail((2400, 2400))
            source = directory / "photo.png"
            image.save(source)
        context = multiprocessing.get_context("spawn")
        for reuse in (False, True):
            parent, child = context.Pipe()
            process = context.Process(
                target=probe, args=(source, str(directory / f"probe-{reuse}.png"), reuse, child)
            )
            process.start()
            child.close()
            assert parent.poll(60)
            report["preview_probe"].append(parent.recv())
            process.join(5)
            assert process.exitcode == 0
            parent.close()
            process.close()
        settings = Settings(workspace=directory / "workspace")
        with closing(Workspace(settings).open()) as workspace:
            uploads = Uploads(workspace, settings)
            jobs = Jobs(workspace, settings, uploads)
            owner = opaque_id()
            for kind in ("pdf", "docx"):
                source = directory / ("preview." + kind)
                if kind == "pdf":
                    source.write_bytes(smoke.minimal_pdf("PDFium local preview benchmark"))
                else:
                    document = Document()
                    document.add_paragraph("Content preview benchmark with real OOXML.")
                    document.add_table(rows=4, cols=3).cell(0, 0).text = "Table cell"
                    document.save(source)
                filename = "source.docx" if kind == "pdf" else "source.pdf"
                identifier = uploads.register(owner, UploadRequest(filename=filename, size_bytes=5))[
                    "upload_id"
                ]
                upload = uploads.items[identifier]
                uploads.path(upload).write_bytes(b"input")
                upload["state"] = "ready"
                workspace.release(identifier)
                request = JobRequest(
                    upload_ids=[identifier],
                    settings={"conversion": "docx_to_pdf" if kind == "pdf" else "pdf_to_docx"},
                )
                job = jobs.items[jobs.create(owner, request)["id"]]
                jobs._inspect(job)
                job["state"] = "completed"
                item = job["files"][0]
                item["state"] = "succeeded"
                target = jobs.output_path(job, item)
                target.write_bytes(source.read_bytes())
                item.update(size_bytes=target.stat().st_size, sha256=file_hash(target))
                with benchmark.MemorySampler(__import__("os").getpid(), workspace.root, 0.05) as sampler:
                    started = time.perf_counter()
                    details, _ = prepare(jobs, job, item)
                    elapsed = time.perf_counter() - started
                started = time.perf_counter()
                cached, _ = prepare(jobs, job, item)
                report["documents"].append(
                    {
                        "kind": kind,
                        "wall_seconds": elapsed,
                        "cached_seconds": time.perf_counter() - started,
                        "cache_hit": cached["cached"],
                        "memory": sampler.result(),
                        "output_preview": details,
                    }
                )
            jobs.close()
    output = args.output
    if output.exists():
        raise RuntimeError("Report already exists")
    output.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
