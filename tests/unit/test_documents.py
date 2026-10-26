"""Real PDF/DOCX library checks; fake LibreOffice tests are explicitly named doubles."""

import hashlib
import os
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from localraw import documents
from localraw.errors import ConversionError
from localraw.models import ConversionSettings


def text_pdf(path, texts=("Editable offline text",), *, encrypted=False, compressed=False):
    """Small generated text PDF, not evidence of OCR or arbitrary layout reconstruction."""
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=612, height=792)
        if text is None:
            continue
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        contents = DecodedStreamObject()
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        contents.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(
            contents.flate_encode() if compressed else contents
        )
    if encrypted:
        writer.encrypt("not-a-supported-password", algorithm="RC4-40")
    with path.open("wb") as stream:
        writer.write(stream)
    writer.close()
    return path


@pytest.fixture
def limits(tmp_path):
    return {
        "document_max_bytes": 64 * 1024**2,
        "document_output_max_bytes": 128 * 1024**2,
        "document_expanded_max_bytes": 256 * 1024**2,
        "document_max_pages": 300,
        "document_max_characters": 2_000_000,
        "pdf_page_content_max_bytes": 16 * 1024**2,
        "document_memory_bytes": 768 * 1024**2,
        "file_timeout_seconds": 30,
        "min_free_bytes": 0,
        "libreoffice": None,
        "scratch_dir": tmp_path / "document-engine-test",
    }


def convert(source, destination, limits, conversion="pdf_to_docx", **callbacks):
    return documents.convert_document(
        source, destination, ConversionSettings(conversion=conversion), limits, **callbacks
    )


def assert_clean(tmp_path, limits):
    assert not limits["scratch_dir"].exists()
    assert not list(tmp_path.rglob("*.part"))


def test_lightweight_discovery_import_does_not_import_document_parsers():
    root = Path(__file__).resolve().parents[2]
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            "import sys; from localraw.documents import libreoffice_executable; "
            "assert not any(n in sys.modules for n in ['docx','pypdf','lxml.etree']); "
            "libreoffice_executable(); print('lightweight')",
        ],
        cwd=root,
        capture_output=True,
        text=True,
        timeout=15,
        check=True,
    )
    assert result.stdout.strip() == "lightweight"


def test_real_text_pdf_to_editable_docx_keeps_pages_and_original(tmp_path, limits):
    source = text_pdf(tmp_path / "original.input", ("First editable page", "Second editable page"))
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    target = tmp_path / "converted.docx"
    stages, admissions = [], []
    result = convert(source, target, limits, progress=stages.append, admit=admissions.append)
    document = Document(target)
    paragraphs = [item.text for item in document.paragraphs]
    assert "First editable page" in paragraphs and "Second editable page" in paragraphs
    assert document.element.xpath("count(//w:br[@w:type='page'])") == 1
    assert result["pages"] == 2 and result["preview"] is None
    assert result["size_bytes"] == target.stat().st_size
    assert result["sha256"] == hashlib.sha256(target.read_bytes()).hexdigest()
    assert any("no OCR" in item for item in result["warnings"])
    assert stages == ["decoding", "encoding", "validating"]
    assert admissions == [limits["document_memory_bytes"]]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert_clean(tmp_path, limits)


def test_mixed_blank_and_text_pages_report_blank_pages(tmp_path, limits):
    source = text_pdf(tmp_path / "input.pdf", (None, "Selectable text", None))
    result = convert(source, tmp_path / "output.docx", limits)
    assert result["pages"] == 3
    assert any("2 page(s)" in warning for warning in result["warnings"])
    assert_clean(tmp_path, limits)


@pytest.mark.parametrize(
    "kind,match",
    [("encrypted", "Encrypted"), ("scanned", "no extractable text"), ("corrupt", "corrupt")],
)
def test_unsupported_pdfs_fail_clearly_without_outputs(tmp_path, limits, kind, match):
    source = tmp_path / "input.pdf"
    if kind == "corrupt":
        source.write_bytes(b"this is not a PDF")
    else:
        text_pdf(source, (None,) if kind == "scanned" else ("Private",), encrypted=kind == "encrypted")
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    target = tmp_path / "output.docx"
    with pytest.raises(ConversionError, match=match):
        convert(source, target, limits)
    assert not target.exists()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert_clean(tmp_path, limits)


@pytest.mark.parametrize(
    "key,value,match",
    [
        ("document_max_pages", 1, "page count"),
        ("document_max_characters", 5, "character limit"),
        ("pdf_page_content_max_bytes", 16, "limit"),
        ("document_output_max_bytes", 1024, "output size"),
        ("document_expanded_max_bytes", 1024, "limit"),
    ],
)
def test_document_resource_limits_leave_no_partial_output(tmp_path, limits, key, value, match):
    source = text_pdf(
        tmp_path / "input.pdf", ("First long enough text", "Second page text"), compressed=True
    )
    limits[key] = value
    target = tmp_path / "output.docx"
    with pytest.raises(ConversionError, match=match):
        convert(source, target, limits)
    assert not target.exists()
    assert_clean(tmp_path, limits)


def test_cancellation_during_encoding_cleans_scratch(tmp_path, limits):
    source = text_pdf(tmp_path / "input.pdf")
    cancelled = False

    def progress(stage):
        nonlocal cancelled
        if stage == "encoding":
            cancelled = True

    target = tmp_path / "output.docx"
    with pytest.raises(ConversionError, match="cancelled"):
        convert(source, target, limits, progress=progress, cancelled=lambda: cancelled)
    assert not target.exists()
    assert_clean(tmp_path, limits)


def test_compressed_pdf_content_is_limited_before_text_parsing(tmp_path, limits):
    source = text_pdf(tmp_path / "compressed.pdf", ("A" * 30_000,), compressed=True)
    assert source.stat().st_size < 1024
    limits["pdf_page_content_max_bytes"] = 1024
    with pytest.raises(ConversionError, match="decoding.*limit"):
        convert(source, tmp_path / "output.docx", limits)
    assert not (tmp_path / "output.docx").exists()
    assert_clean(tmp_path, limits)


def test_owned_docx_writer_preserves_python_docx_package_serialization(tmp_path, limits):
    document = Document()
    document.add_heading("Editable output", 0)
    document.add_paragraph("Paragraph with a tab\there and a line\nbreak.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Left", "Right"
    expected, actual = tmp_path / "expected.docx", tmp_path / "actual.docx"
    document.save(expected)
    documents._save_docx(document, actual, limits, lambda: False)
    with zipfile.ZipFile(expected) as first, zipfile.ZipFile(actual) as second:
        assert first.namelist() == second.namelist()
        assert all(first.read(name) == second.read(name) for name in first.namelist())


def test_docx_split_external_field_is_rejected_before_native_engine(tmp_path, limits, monkeypatch):
    from docx.oxml import OxmlElement

    source = tmp_path / "input.docx"
    document = Document()
    paragraph = document.add_paragraph()
    for value in (" D", "DEAUTO remote application"):
        instruction = OxmlElement("w:instrText")
        instruction.text = value
        paragraph.add_run().element.append(instruction)
    document.save(source)
    monkeypatch.setattr(documents, "libreoffice_executable", lambda configured=None: "engine-double")
    monkeypatch.setattr(
        documents.subprocess,
        "Popen",
        lambda *args, **kwargs: pytest.fail("External field reached native engine"),
    )
    with pytest.raises(ConversionError, match="external data fields"):
        convert(source, tmp_path / "output.pdf", limits, "docx_to_pdf")
    assert_clean(tmp_path, limits)


def test_xml_content_type_with_non_xml_extension_still_gets_entity_validation(
    tmp_path, limits, monkeypatch
):
    source, original = tmp_path / "input.docx", tmp_path / "original.docx"
    document = Document()
    document.add_paragraph("Safe text")
    document.save(original)
    with (
        zipfile.ZipFile(original) as first,
        zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as second,
    ):
        for entry in first.infolist():
            data = first.read(entry)
            if entry.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/word/payload.dat" ContentType="application/xml"/></Types>',
                )
            second.writestr(entry, data)
        second.writestr(
            "word/payload.dat",
            b'<!DOCTYPE x [<!ENTITY e SYSTEM "https://invalid.example/never-fetch">]><x>&e;</x>',
        )
    monkeypatch.setattr(documents, "libreoffice_executable", lambda configured=None: "engine-double")
    monkeypatch.setattr(
        documents.subprocess,
        "Popen",
        lambda *args, **kwargs: pytest.fail("XML entity reached native engine"),
    )
    with pytest.raises(ConversionError, match="entity"):
        convert(source, tmp_path / "output.pdf", limits, "docx_to_pdf")
    assert_clean(tmp_path, limits)


def test_missing_libreoffice_has_actionable_error(tmp_path, limits, monkeypatch):
    source = tmp_path / "input.docx"
    document = Document()
    document.add_paragraph("Offline document")
    document.save(source)
    monkeypatch.setattr(documents, "libreoffice_executable", lambda configured=None: None)
    target = tmp_path / "output.pdf"
    with pytest.raises(ConversionError, match="LibreOffice installed locally"):
        convert(source, target, limits, "docx_to_pdf")
    assert not target.exists()
    assert_clean(tmp_path, limits)


@pytest.mark.parametrize(
    "name,data,match",
    [
        ("../../escape.xml", b"<x/>", "corrupt|unsafe"),
        ("word/vbaProject.bin", b"macro double", "macros"),
        ("word/activeX/activeX1.xml", b"<x/>", "active content"),
        ("word/embeddings/oleObject1.bin", b"embedded double", "embedded"),
        (
            "word/vector.dat",
            b'<svg xmlns="http://www.w3.org/2000/svg"><image href="https://invalid.example/never-fetch"/></svg>',
            "SVG resources",
        ),
        ("word/evil.xml", b'<!DOCTYPE x [<!ENTITY x SYSTEM "file:///private">]><x>&x;</x>', "entity"),
        (
            "word/_rels/evil.xml.rels",
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="evil" Type="image" Target="https://invalid.example/resource" TargetMode="External"/></Relationships>',
            "external links",
        ),
        (
            "word/_rels/evil.xml.rels",
            b'<Relationships><Relationship Target="../../../../outside"/></Relationships>',
            "unsafe internal",
        ),
    ],
)
def test_unsafe_docx_packages_rejected_before_libreoffice(
    tmp_path, limits, name, data, match, monkeypatch
):
    source = tmp_path / "input.docx"
    document = Document()
    document.add_paragraph("Safe original text")
    document.save(source)
    with zipfile.ZipFile(source, "a") as package:
        package.writestr(name, data)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    monkeypatch.setattr(
        documents, "libreoffice_executable", lambda configured=None: "explicit-engine-double"
    )
    monkeypatch.setattr(
        documents.subprocess,
        "Popen",
        lambda *args, **kwargs: pytest.fail("Unsafe input reached LibreOffice"),
    )
    with pytest.raises(ConversionError, match=match):
        convert(source, tmp_path / "output.pdf", limits, "docx_to_pdf")
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert_clean(tmp_path, limits)


def test_existing_output_and_unsafe_scratch_are_preserved(tmp_path, limits):
    source = text_pdf(tmp_path / "input.pdf")
    target = tmp_path / "output.docx"
    target.write_bytes(b"existing user output")
    with pytest.raises(ConversionError, match="already exists"):
        convert(source, target, limits)
    assert target.read_bytes() == b"existing user output"
    limits["scratch_dir"] = tmp_path.parent / "document-outside"
    with pytest.raises(ConversionError, match="staging location"):
        convert(source, tmp_path / "other.docx", limits)
    assert not limits["scratch_dir"].exists()


def test_libreoffice_subprocess_double_uses_private_profile_and_validated_output(
    tmp_path, limits, monkeypatch
):
    """Tests the process contract only; this is not native LibreOffice conversion evidence."""
    source = tmp_path / "input.docx"
    document = Document()
    document.add_paragraph("Native engine contract")
    document.save(source)
    recorded = []

    class ExplicitLibreOfficeDouble:
        pid = -1
        returncode = 0

        def __init__(self, arguments, **kwargs):
            assert kwargs["cwd"].samefile(limits["scratch_dir"])
            recorded.append((arguments, kwargs))
            output = Path(arguments[arguments.index("--outdir") + 1]) / "input.pdf"
            text_pdf(output, ("Explicit subprocess double",))

        def poll(self):
            return 0

        def wait(self, timeout=None):
            return 0

    monkeypatch.setattr(
        documents, "libreoffice_executable", lambda configured=None: "explicit-engine-double"
    )
    monkeypatch.setattr(documents.subprocess, "Popen", ExplicitLibreOfficeDouble)
    monkeypatch.setattr(documents, "_stop_process_tree", lambda process: None)
    target = tmp_path / "output.pdf"
    result = convert(source, target, limits, "docx_to_pdf")
    assert result["pages"] == 1
    with target.open("rb") as stream:
        assert "Explicit subprocess double" in PdfReader(stream).pages[0].extract_text()
    arguments, options = recorded[0]
    assert "--headless" in arguments and "--norestore" in arguments
    assert any(value.startswith("-env:UserInstallation=file:") for value in arguments)
    assert options["stdout"] == subprocess.DEVNULL and options["stderr"] == subprocess.DEVNULL
    assert_clean(tmp_path, limits)


@pytest.mark.skipif(os.name != "nt", reason="Windows native profile-path constraint")
def test_long_office_profile_without_short_name_has_actionable_error(tmp_path, monkeypatch):
    monkeypatch.setattr(documents, "_office_cli_path", lambda path: Path(path))
    profile = tmp_path / ("long-workspace-" * 8) / "profile"
    with pytest.raises(ConversionError, match="shorter workspace"):
        documents._office_profile_path(profile)


@pytest.mark.skipif(
    not documents.libreoffice_executable(os.environ.get("LOCALRAW_TEST_LIBREOFFICE")),
    reason="Local LibreOffice is unavailable; set LOCALRAW_TEST_LIBREOFFICE for native validation",
)
def test_installed_libreoffice_real_docx_to_pdf(tmp_path, limits):
    limits["libreoffice"] = os.environ.get("LOCALRAW_TEST_LIBREOFFICE")
    limits["file_timeout_seconds"] = 120  # Allow a cold first launch/font scan in a fresh profile.
    source = tmp_path / "input.docx"
    document = Document()
    document.add_heading("Offline document", 0)
    document.add_paragraph("Editable local source converted by LibreOffice.")
    document.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    target = tmp_path / "output.pdf"
    result = convert(source, target, limits, "docx_to_pdf")
    with target.open("rb") as stream:
        reader = PdfReader(stream)
        assert "Editable local source" in reader.pages[0].extract_text()
    assert result["pages"] >= 1 and result["sha256"] == hashlib.sha256(target.read_bytes()).hexdigest()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert_clean(tmp_path, limits)
