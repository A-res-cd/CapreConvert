"""Real document-engine API checks with programmatically authored, valid files."""

import os
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from localraw.app import create_app
from tests.support import login, upload, wait_job


def text_pdf():
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (CapreConvert editable text) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    result = BytesIO()
    writer.write(result)
    writer.close()
    return result.getvalue()


def test_real_pdf_docx_api_download_restart_and_ownership(settings, ample_memory):
    base = "http://127.0.0.1:8765"
    app = create_app(settings)
    with TestClient(app, base_url=base, client=("127.0.0.1", 40000), headers={"Origin": base}) as client:
        login(client, app)
        raw = upload(client, name="text.pdf", content=text_pdf())
        job = client.post(
            "/api/v1/jobs", json={"upload_ids": [raw], "settings": {"conversion": "pdf_to_docx"}}
        ).json()
        final = wait_job(client, job["id"])
        assert final["state"] == "completed", client.get(f"/api/v1/jobs/{job['id']}/errors").text
        item = client.get(f"/api/v1/jobs/{job['id']}/files").json()["items"][0]
        assert item["output_name"] == "files/text.docx" and not item.get("preview")
        route = f"/api/v1/jobs/{job['id']}/results/{item['output_id']}"
        response = client.get(route)
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
        payload = response.content
        document = Document(BytesIO(payload))
        assert "CapreConvert editable text" in "\n".join(p.text for p in document.paragraphs)
        assert client.get(f"/api/v1/jobs/{job['id']}/previews/{item['output_id']}").status_code == 404
        cookies = dict(client.cookies)
        client.cookies.clear()
        assert client.get(route).status_code == 401
    restored = create_app(settings, start_scheduler=False)
    with TestClient(restored, base_url=base, client=("127.0.0.1", 40000), cookies=cookies) as client:
        assert client.get(route).content == payload
        assert client.get(f"/api/v1/jobs/{job['id']}/archive").status_code == 200


def test_real_libreoffice_docx_pdf_api(settings, ample_memory):
    engine = os.environ.get("LOCALRAW_TEST_LIBREOFFICE")
    if not engine:
        pytest.skip("External LibreOffice path required for real DOCX-to-PDF API test")
    settings.libreoffice = Path(engine)
    source = BytesIO()
    document = Document()
    document.add_heading("CapreConvert local document", level=1)
    document.add_paragraph("This document was converted with local LibreOffice.")
    document.save(source)
    app = create_app(settings)
    base = "http://127.0.0.1:8765"
    with TestClient(app, base_url=base, client=("127.0.0.1", 40000), headers={"Origin": base}) as client:
        login(client, app)
        descriptor = next(
            item
            for item in client.get("/api/v1/capabilities").json()["conversions"]
            if item["id"] == "docx_to_pdf"
        )
        assert descriptor["available"] and not descriptor["setup_message"]
        raw = upload(client, name="letter.docx", content=source.getvalue())
        job = client.post(
            "/api/v1/jobs", json={"upload_ids": [raw], "settings": {"conversion": "docx_to_pdf"}}
        ).json()
        final = wait_job(client, job["id"], timeout=120)
        assert final["state"] == "completed", client.get(f"/api/v1/jobs/{job['id']}/errors").text
        item = client.get(f"/api/v1/jobs/{job['id']}/files").json()["items"][0]
        response = client.get(f"/api/v1/jobs/{job['id']}/results/{item['output_id']}")
        assert response.headers["content-type"] == "application/pdf"
        reader = PdfReader(BytesIO(response.content))
        assert len(reader.pages) == 1
        assert "CapreConvert local document" in reader.pages[0].extract_text()
        reader.close()
        assert not list(app.state.workspace.root.rglob("document-*"))
        assert not app.state.workspace.reservations
