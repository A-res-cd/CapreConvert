"""Actual raster API workers and actual PDFium/OOXML preview workers."""

import concurrent.futures
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from localraw.app import create_app
from localraw.storage import file_hash
from tests.integration.test_document_api import text_pdf
from tests.support import login, upload, wait_job

BASE = "http://127.0.0.1:8765"


def encoded(kind="PNG", alpha=True):
    result = BytesIO()
    with Image.new(
        "RGBA" if alpha else "RGB", (60, 40), (20, 70, 120, 40) if alpha else (20, 70, 120)
    ) as image:
        image.save(result, format=kind)
    return result.getvalue()


def test_real_raster_worker_exports_alpha_restart_ownership(settings, ample_memory):
    app = create_app(settings)
    with TestClient(app, base_url=BASE, client=("127.0.0.1", 40000), headers={"Origin": BASE}) as client:
        login(client, app)
        ids = [upload(client, name=name, content=encoded()) for name in ("one.png", "two.png")]
        response = client.post(
            "/api/v1/jobs",
            json={
                "upload_ids": ids,
                "settings": {
                    "conversion": "png_to_webp",
                    "raster": {"webp_lossless": True, "max_width": 30},
                },
            },
        )
        assert response.status_code == 202, response.text
        job = response.json()
        assert wait_job(client, job["id"])["state"] == "completed"
        prefix = f"/api/v1/jobs/{job['id']}"
        files = client.get(prefix + "/files").json()["items"]
        result = client.get(prefix + "/results/" + files[0]["output_id"])
        assert result.headers["content-type"] == "image/webp"
        assert int(result.headers["content-length"]) == len(result.content)
        with Image.open(BytesIO(result.content)) as output:
            assert output.size == (30, 20) and output.mode == "RGBA"
        for item in files:
            preview = client.get(prefix + "/previews/" + item["output_id"])
            assert preview.headers["content-type"] == "image/png"
            with Image.open(BytesIO(preview.content)) as image:
                assert image.mode == "RGBA"
        selected = client.post(
            prefix + "/exports", json={"output_ids": [item["output_id"] for item in files]}
        ).json()
        with ZipFile(
            BytesIO(client.get(prefix + "/exports/" + selected["export_id"]).content)
        ) as archive:
            assert len(archive.namelist()) == 2
            assert archive.read(files[0]["output_name"]) == result.content
        # ZIP endpoint remains ZIP even for one selection.
        single = client.post(prefix + "/exports", json={"output_ids": [files[0]["output_id"]]}).json()
        assert client.get(prefix + "/exports/" + single["export_id"]).content[:2] == b"PK"
        cookies = dict(client.cookies)
        client.cookies.clear()
        assert client.get(prefix + "/previews/" + files[0]["output_id"]).status_code == 401
        assert client.get(prefix + "/results/" + files[0]["output_id"]).status_code == 401
    with TestClient(
        create_app(settings, start_scheduler=False),
        base_url=BASE,
        client=("127.0.0.1", 40000),
        cookies=cookies,
    ) as restored:
        assert restored.get(prefix + "/results/" + files[0]["output_id"]).content == result.content
        assert restored.get(prefix + "/previews/" + files[0]["output_id"]).status_code == 200


def seeded_preview(app, client, kind, payload):
    source = upload(
        client, name="source.docx" if kind == "pdf" else "source.pdf", content=b"retained source"
    )
    conversion = "docx_to_pdf" if kind == "pdf" else "pdf_to_docx"
    response = client.post(
        "/api/v1/jobs", json={"upload_ids": [source], "settings": {"conversion": conversion}}
    )
    job = app.state.jobs.items[response.json()["id"]]
    app.state.jobs._inspect(job)
    job["state"] = "completed"
    item = job["files"][0]
    item["state"] = "succeeded"
    path = app.state.jobs.output_path(job, item)
    path.write_bytes(payload)
    item.update(size_bytes=len(payload), sha256=file_hash(path))
    app.state.jobs.save(job)
    return job, item, f"/api/v1/jobs/{job['id']}"


def test_pdf_lazy_preview_cache_dedup_limits_download_independence(settings, ample_memory):
    app = create_app(settings, start_scheduler=False)
    with TestClient(app, base_url=BASE, client=("127.0.0.1", 40000), headers={"Origin": BASE}) as client:
        login(client, app)
        job, item, prefix = seeded_preview(app, client, "pdf", text_pdf())
        route = prefix + "/output-preview/" + item["output_id"]
        assert not list(app.state.jobs.directory(job).glob("pv-*"))
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            initial = list(pool.map(lambda _: client.get(route), range(2)))
        assert sorted(result.headers["x-preview-cached"] for result in initial) == ["false", "true"]
        first = next(result for result in initial if result.headers["x-preview-cached"] == "false")
        assert first.status_code == 200, first.text
        assert first.headers["x-preview-pages"] == "1" and first.headers["x-preview-cached"] == "false"
        with Image.open(BytesIO(first.content)) as image:
            assert image.width * image.height <= settings.preview_max_pixels
            assert image.getextrema()[0][0] < 255  # rendered text, not a blank placeholder
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: client.get(route), range(2)))
        assert all(result.headers["x-preview-cached"] == "true" for result in results)
        assert client.get(route + "?page=2").status_code == 422
        assert client.get(route + "?page=0").status_code == 422
        assert client.get(prefix + "/results/" + item["output_id"]).content == text_pdf()
        # Invalid cache forces bounded regeneration; failure leaves actual output downloadable.
        for path in app.state.jobs.directory(job).glob("pv-*"):
            path.unlink()
        settings.preview_max_bytes = 10
        assert client.get(route).status_code == 422
        assert client.get(prefix + "/results/" + item["output_id"]).status_code == 200
        client.cookies.clear()
        assert client.get(route).status_code == 401


def test_docx_preview_safe_bounded_tables_no_office(settings, ample_memory):
    settings.libreoffice = None
    app = create_app(settings, start_scheduler=False)
    source = BytesIO()
    document = Document()
    document.add_paragraph('<script>alert("unsafe")</script>')
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Basic cell"
    document.save(source)
    with TestClient(app, base_url=BASE, client=("127.0.0.1", 40000), headers={"Origin": BASE}) as client:
        login(client, app)
        job, item, prefix = seeded_preview(app, client, "docx", source.getvalue())
        route = prefix + "/output-preview/" + item["output_id"]
        response = client.get(route)
        assert response.status_code == 200, response.text
        content = response.json()
        assert content["blocks"][0]["text"].startswith("<script>")
        assert content["blocks"][1]["rows"][0][0] == "Basic cell"
        assert "not exact" in content["notice"]
        settings.preview_text_characters = 5
        response = client.get(route)
        assert response.json()["truncated"]
        assert len(response.json()["blocks"][0]["text"]) == 5
        assert client.get(prefix + "/results/" + item["output_id"]).content == source.getvalue()


def test_preview_serializes_with_conversion_resource_lock_and_quota(settings, ample_memory):
    import threading
    import time

    app = create_app(settings, start_scheduler=False)
    with TestClient(app, base_url=BASE, client=("127.0.0.1", 40000), headers={"Origin": BASE}) as client:
        login(client, app)
        job, item, prefix = seeded_preview(app, client, "pdf", text_pdf())
        route = prefix + "/output-preview/" + item["output_id"]
        app.state.jobs.execution_lock.acquire()
        result = []
        thread = threading.Thread(target=lambda: result.append(client.get(route)))
        thread.start()
        try:
            time.sleep(0.15)
            assert not result
            assert not list(app.state.jobs.directory(job).glob("pv-*"))
        finally:
            app.state.jobs.execution_lock.release()
        thread.join(timeout=15)
        assert result[0].status_code == 200
        assert not app.state.workspace.reservations
        for path in app.state.jobs.directory(job).glob("pv-*"):
            path.unlink()
        settings.workspace_quota_bytes = app.state.workspace.used_bytes() + 100
        assert client.get(route).status_code == 507
        assert client.get(prefix + "/results/" + item["output_id"]).status_code == 200


def test_raster_alias_actual_mismatch_and_zip_policy(settings, ample_memory):
    app = create_app(settings)
    with TestClient(app, base_url=BASE, client=("127.0.0.1", 40000), headers={"Origin": BASE}) as client:
        login(client, app)
        source = upload(client, name="alias.jpeg", content=encoded("JPEG", False))
        job = client.post(
            "/api/v1/jobs", json={"upload_ids": [source], "settings": {"conversion": "jpg_to_png"}}
        ).json()
        assert wait_job(client, job["id"])["state"] == "completed"
        wrong = upload(client, name="mismatch.jpg", content=encoded())
        job = client.post(
            "/api/v1/jobs", json={"upload_ids": [wrong], "settings": {"conversion": "jpg_to_png"}}
        ).json()
        assert wait_job(client, job["id"])["state"] == "failed"
        assert "Actual image format" in client.get(f"/api/v1/jobs/{job['id']}/errors").text
        archive = upload(client, name="unsupported.zip", content=b"not ZIP")
        assert (
            client.post(
                "/api/v1/jobs", json={"upload_ids": [archive], "settings": {"conversion": "png_to_png"}}
            ).status_code
            == 422
        )


def test_preview_cache_survives_restart_and_corruption_recovers(settings, ample_memory):
    app = create_app(settings, start_scheduler=False)
    with TestClient(app, base_url=BASE, client=("127.0.0.1", 40000), headers={"Origin": BASE}) as client:
        login(client, app)
        job, item, prefix = seeded_preview(app, client, "pdf", text_pdf())
        route = prefix + "/output-preview/" + item["output_id"]
        assert client.get(route).status_code == 200
        cookies = dict(client.cookies)
    restored = create_app(settings, start_scheduler=False)
    with TestClient(restored, base_url=BASE, client=("127.0.0.1", 40000), cookies=cookies) as client:
        response = client.get(route)
        assert response.headers["x-preview-cached"] == "true"
        sidecar = next(restored.state.jobs.directory(job).glob("pv-*.meta"))
        sidecar.write_text("[]", encoding="utf-8")
        response = client.get(route)
        assert response.status_code == 200 and response.headers["x-preview-cached"] == "false"
        assert not list(restored.state.workspace.root.rglob("*.part"))
